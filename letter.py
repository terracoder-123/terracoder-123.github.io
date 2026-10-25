#!/usr/bin/env python3
"""Build a cover letter .docx in the same layout as Althaf's own letters.

  python3 letter.py N spec.json

N is the shortlist number from `hunt.py search`. spec.json holds what is
specific to the job:

  {
    "recipient": ["Hiring Manager", "Enerven", "Hindmarsh SA"],
    "re": "Application for Data Analyst",
    "salutation": "Dear Hiring Manager,",
    "paragraphs": ["opening ...", "skills ...", "project evidence ..."],
    "closing": "optional; defaults to the standard closing below",
    "city": "optional override, e.g. Wollongong",
    "relocation": "optional override: adelaide | south_australia | other | home | remote"
  }

The relocation / working-rights paragraph is added automatically from
profile.json using the job's location, so every letter outside Melbourne
says Althaf will move for the role. The phone number is read from the
APPLICANT_PHONE environment variable so it never lands in the public repo.
"""

import datetime as dt
import json
import os
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

import hunt

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x40, 0x40, 0x40)
CLOSING = ("My resume is attached for your consideration. I would welcome the opportunity to discuss "
           "how I could contribute to {company}. Thank you for your time and consideration.")
# Claims the profile says not to make; a hit stops the build so the letter can be fixed.
RED_FLAGS = [r"\b\d+\+?\s*(years|yrs)('|’)? (of )?(professional |commercial |industry )?experience",
             r"\b(i am|i'm) an australian citizen", r"\bpermanent resident", r"\bsecurity clearance\b",
             r"own expense", r"\b(experience|experienced|proficient) (in|with) (power bi|tableau|azure|aws|snowflake|databricks|dbt)"]


def para(doc, text, *, size=11, bold=False, color=None, after=0, line=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Twips(after)
    if line:
        pf.line_spacing = line / 240
    p.alignment = align
    r = p.add_run(text)
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    r.font.size = Pt(size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = color
    return p


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:color": "1F3864", "w:sz": "8", "w:space": "6"}.items():
        b.set(qn(k), v)
    bdr.append(b)
    pPr.insert(0, bdr)


def relocation_paragraph(job):
    templates = hunt.PROFILE["relocation"]
    return templates[job["relocation"]].format(city=job.get("city") or "your location")


def check(text):
    problems = [m.group(0) for rx in RED_FLAGS for m in re.finditer(rx, text, re.I)]
    words = len(text.split())
    if words > 480:
        problems.append(f"{words} words; keep it to one page (about 450)")
    return problems


def build(job, spec):
    a = hunt.PROFILE["applicant"]
    phone = os.environ.get(a["phone_env"], "").strip()
    contact = "  |  ".join(x for x in [a["suburb"], phone, a["email"], a["linkedin"]] if x)

    body = list(spec["paragraphs"]) + [relocation_paragraph(job),
                                       spec.get("closing") or CLOSING.format(company=job["company"])]
    problems = check(" ".join(body))
    if problems:
        sys.exit("Letter not written, fix these first:\n  - " + "\n  - ".join(problems))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)
    sec.top_margin = sec.bottom_margin = Twips(1000)
    sec.left_margin = sec.right_margin = Twips(1200)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    doc.core_properties.author = a["name"]
    doc.core_properties.title = f"Cover letter - {job['title']} - {job['company']}"

    para(doc, a["name"], size=18, bold=True, color=NAVY)
    bottom_border(para(doc, contact, size=9.5, color=GREY, after=240))
    para(doc, dt.date.fromisoformat(hunt.today()).strftime("%-d %B %Y"), after=200)
    recipient = spec.get("recipient") or ["Hiring Manager", job["company"], job["location"]]
    for i, line in enumerate(recipient):
        para(doc, line, after=200 if i == len(recipient) - 1 else 0)
    para(doc, spec.get("re") or f"Re: Application for {job['title']}", bold=True, after=160, line=264)
    para(doc, spec.get("salutation") or "Dear Hiring Manager,", after=160, line=264)
    for text in body:
        para(doc, text, after=160, line=264, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(doc, "Yours sincerely,", after=480)
    para(doc, a["name"], bold=True)

    slug = re.sub(r"[^A-Za-z0-9]+", "_", f"{job['company']}_{job['title']}").strip("_")[:60]
    path = hunt.out_dir() / f"{job['n']:02d}_Althaf_Rismy_Cover_Letter_{slug}.docx"
    doc.save(path)
    return path


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    n = int(sys.argv[1])
    spec = json.loads(Path(sys.argv[2]).read_text())
    job, _ = hunt.shortlist_item(n)
    job = {**job, **{k: spec[k] for k in ("city", "relocation") if k in spec}}
    # A false positive (e.g. "citizens, permanent residents and anyone with work rights") can be
    # cleared by putting "eligibility_checked": "<why Althaf is eligible>" in the spec.
    if ("clearance/citizenship" in job["flags"] and not hunt.PROFILE["eligibility"]["australian_citizen"]
            and not spec.get("eligibility_checked")):
        sys.exit(f"#{n} needs clearance or citizenship; skip it with hunt.py skip {n} \"needs citizenship/clearance\"")
    if not os.environ.get(hunt.PROFILE["applicant"]["phone_env"]):
        print("warning: APPLICANT_PHONE not set, letter will have no phone number", file=sys.stderr)
    path = build(job, spec)

    apps_file = hunt.run_dir() / "applications.json"
    apps = [x for x in hunt.load_json(apps_file, []) if x["n"] != n]
    apps.append({"n": n, "status": "letter", "file": path.name, "key": job["key"]})
    hunt.save_json(apps_file, apps)
    print(f"Wrote {path.relative_to(hunt.ROOT)}")


if __name__ == "__main__":
    main()
